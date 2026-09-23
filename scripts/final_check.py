import os
import docx
import json
import subprocess

files = [
    'Likhita_FoodWise_AI.ipynb',
    'requirements.txt',
    'Likhita_FoodWise_AI_ProjectReport.docx',
    'README.md',
    'data/foodwise_historical_data.csv'
]

print("=== 1. FILE EXISTENCE & SIZE CHECK ===")
for f in files:
    full_p = os.path.join(r'd:\FoodWise AI', f)
    exists = os.path.exists(full_p)
    size = os.path.getsize(full_p) if exists else 0
    status = "PASS" if exists else "FAIL"
    print(f"[{status}] {f:40s} | Size: {size:,} bytes")
    assert exists, f"Missing file: {f}"

print("\n=== 2. NOTEBOOK STRUCTURE CHECK ===")
with open(r'd:\FoodWise AI\Likhita_FoodWise_AI.ipynb', 'r', encoding='utf-8') as f:
    nb_data = json.load(f)
cells_count = len(nb_data['cells'])
print(f"Total Notebook Cells: {cells_count}")
assert cells_count == 31, f"Expected 31 cells, got {cells_count}"

print("\n=== 3. REPORT STRUCTURE CHECK ===")
doc = docx.Document(r'd:\FoodWise AI\Likhita_FoodWise_AI_ProjectReport.docx')
print(f"Report Paragraphs: {len(doc.paragraphs)} | Tables: {len(doc.tables)}")

print("\n=== 4. WEB APPLICATION BUILD CHECK ===")
res_type = subprocess.run(["npm.cmd", "run", "typecheck"], capture_output=True, text=True, cwd=r'd:\FoodWise AI')
print("Typecheck Exit Code:", res_type.returncode)
assert res_type.returncode == 0, "Typecheck failed!"

res_build = subprocess.run(["npm.cmd", "run", "build"], capture_output=True, text=True, cwd=r'd:\FoodWise AI')
print("Build Exit Code:", res_build.returncode)
assert res_build.returncode == 0, "Build failed!"

print("\n==================================================")
print("ALL FINAL VALIDATION CHECKS PASSED SUCCESSFULLY!")
print("==================================================")
